"""Generate deterministic TEST FIXTURE resume import samples.

All identity, organizations, dates, and metrics come from the explicitly
fictional JSON fixture and do not correspond to a real person.
"""

from __future__ import annotations

import json
import textwrap
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "tests" / "fixtures" / "resume-import-fixture.json"
OUTPUT = ROOT / "tmp" / "qa" / "resume-import"
SECTIONS = [
    ("个人简介", "summary"),
    ("教育经历", "education"),
    ("实习经历", "experience"),
    ("项目经历", "project"),
    ("校园经历", "campus"),
    ("技能特长", "skills"),
    ("证书荣誉", "certificate"),
    ("自我评价", "evaluation"),
    ("作品链接", "portfolio"),
]


def load_fixture() -> dict:
    return json.loads(SOURCE.read_text(encoding="utf-8"))


def canonical_lines(data: dict) -> list[str]:
    lines = [
        data["fixtureLabel"],
        f"姓名：{data['name']}",
        f"邮箱：{data['email']}",
        f"手机：{data['phone']}",
        f"所在地：{data['city']}",
        f"求职目标：{data['target']}",
        "",
    ]
    for heading, key in SECTIONS:
        lines.append(heading)
        value = data[key]
        entries = value if isinstance(value, list) else [value]
        for index, entry in enumerate(entries):
            lines.extend(entry.splitlines())
            if index < len(entries) - 1:
                lines.append("")
        lines.append("")
    return lines


def set_run_font(run, size: float, bold: bool = False, color="25302A") -> None:
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def create_docx(data: dict, lines: list[str]) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.right_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(18)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    heading_names = {heading for heading, _ in SECTIONS}
    for index, line in enumerate(lines):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.15
        if index == 0:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(paragraph.add_run(line), 9, True, "9B1C1C")
        elif line == f"姓名：{data['name']}":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(paragraph.add_run(line), 20, True, "173D2D")
            paragraph.paragraph_format.space_after = Pt(6)
        elif line in heading_names:
            paragraph.style = document.styles["Heading 1"]
            paragraph.paragraph_format.space_before = Pt(9)
            paragraph.paragraph_format.space_after = Pt(4)
            set_run_font(paragraph.add_run(line), 12, True, "225E45")
        else:
            set_run_font(paragraph.add_run(line), 10.5)
    document.core_properties.title = "TEST FIXTURE resume import sample"
    document.core_properties.author = "TEST FIXTURE generator"
    document.save(OUTPUT / "resume-fixture.docx")


def create_pdf(lines: list[str]) -> None:
    pdfmetrics.registerFont(
        TTFont("MicrosoftYaHei", r"C:\Windows\Fonts\msyh.ttc", subfontIndex=0)
    )
    canvas = Canvas(str(OUTPUT / "resume-fixture.pdf"), pagesize=A4)
    width, height = A4
    y = height - 42
    for line in lines:
        wrapped = textwrap.wrap(line, width=72) or [""]
        for item in wrapped:
            if y < 42:
                canvas.showPage()
                y = height - 42
            canvas.setFont("MicrosoftYaHei", 8.5)
            canvas.drawString(42, y, item)
            y -= 10.5
        if not line:
            y -= 2
    canvas.save()


def find_font(size: int) -> ImageFont.FreeTypeFont:
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    raise FileNotFoundError("No Chinese font found for QA fixture images")


def create_images(lines: list[str]) -> None:
    image = Image.new("RGB", (1800, 2600), "white")
    draw = ImageDraw.Draw(image)
    body = find_font(30)
    heading = find_font(34)
    y = 45
    heading_names = {heading_name for heading_name, _ in SECTIONS}
    for line in lines:
        font = heading if line in heading_names else body
        wrapped = textwrap.wrap(line, width=46) or [""]
        for item in wrapped:
            if y > 2540:
                raise ValueError("Fixture image content exceeds one page")
            draw.text((55, y), item, font=font, fill="black")
            y += 42 if item else 22
    image.save(OUTPUT / "resume-fixture.png", optimize=True)
    image.save(OUTPUT / "resume-fixture.jpg", quality=94, subsampling=0)


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    data = load_fixture()
    lines = canonical_lines(data)
    text = "\n".join(lines).rstrip() + "\n"
    (OUTPUT / "resume-fixture.txt").write_text(text, encoding="utf-8")
    (OUTPUT / "golden-text.txt").write_text(text, encoding="utf-8")
    create_docx(data, lines)
    create_pdf(lines)
    create_images(lines)
    print(f"Generated TEST FIXTURE samples in {OUTPUT}")


if __name__ == "__main__":
    main()
